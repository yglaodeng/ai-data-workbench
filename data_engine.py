from __future__ import annotations

import base64
import hashlib
import io
import json
import math
import os
import re
import shutil
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_FILES = 5
MAX_ROWS = 100_000
MAX_COLUMNS = 1_000
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
PREVIEW_ROWS = 20
SUPPORTED_SUFFIXES = {".xlsx", ".csv", ".tsv", ".json"}
OUTPUT_MODES = {"linked", "standalone"}


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _json_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if pd.isna(value):
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if hasattr(value, "item"):
        return value.item()
    return value


def frame_records(frame: pd.DataFrame, limit: int = PREVIEW_ROWS) -> list[dict[str, Any]]:
    return [
        {str(column): _json_value(value) for column, value in row.items()}
        for row in frame.head(limit).to_dict(orient="records")
    ]


def unique_columns(columns: list[Any]) -> list[str]:
    result: list[str] = []
    seen: dict[str, int] = {}
    for index, raw in enumerate(columns, start=1):
        base = str(raw).strip() or f"未命名字段{index}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        result.append(base if count == 0 else f"{base}_{count + 1}")
    return result


def field_role(column: str, series: pd.Series | None = None) -> str:
    name = str(column).strip().lower()
    if re.search(r"(^id$|编号|编码|单号|票号|订单号|流水号|sku|pnr)", name):
        return "identifier"
    if re.search(r"日期|时间|年月|月份|季度|年度", name):
        return "date"
    if re.search(r"率|比例|折扣|扣率|占比", name):
        return "percentage"
    if re.search(r"金额|售价|价格|收入|成本|利润|费用|服务费|底价|税", name):
        return "amount"
    if re.search(r"数量|销量|件数|人数|次数|订单数", name):
        return "quantity"
    if series is not None and pd.api.types.is_datetime64_any_dtype(series):
        return "date"
    if series is not None and pd.api.types.is_bool_dtype(series):
        return "boolean"
    if series is not None and pd.api.types.is_numeric_dtype(series):
        return "number"
    return "category"


def profile_frame(frame: pd.DataFrame) -> dict[str, Any]:
    fields = []
    for column in frame.columns:
        series = frame[column]
        non_empty = series.dropna()
        if pd.api.types.is_datetime64_any_dtype(series):
            inferred = "date"
        elif pd.api.types.is_numeric_dtype(series):
            inferred = "number"
        elif pd.api.types.is_bool_dtype(series):
            inferred = "boolean"
        else:
            inferred = "text"
        fields.append({
            "name": str(column),
            "type": inferred,
            "role": field_role(str(column), series),
            "emptyCount": int(series.isna().sum()),
            "nonEmptyCount": int(series.notna().sum()),
            "uniqueCount": int(non_empty.astype(str).nunique()),
        })
    return {
        "rowCount": int(len(frame)),
        "columnCount": int(len(frame.columns)),
        "columns": [str(column) for column in frame.columns],
        "fields": fields,
    }


def summary_rows(frame: pd.DataFrame) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for column in frame.columns:
        role = field_role(str(column), frame[column])
        if role not in {"amount", "quantity", "percentage", "number"}:
            continue
        values = pd.to_numeric(frame[column], errors="coerce").dropna()
        if values.empty:
            continue
        rows.append({
            "field": str(column),
            "role": role,
            "count": int(values.count()),
            "sum": float(values.sum()),
            "average": float(values.mean()),
            "min": float(values.min()),
            "max": float(values.max()),
        })
    return rows


def _excel_value(value: Any) -> Any:
    value = _json_value(value)
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def _write_frame_sheet(workbook: Any, title: str, frame: pd.DataFrame) -> Any:
    if title in workbook.sheetnames:
        del workbook[title]
    sheet = workbook.create_sheet(title)
    header_fill = PatternFill("solid", fgColor="3157D5")
    for column_index, column in enumerate(frame.columns, start=1):
        cell = sheet.cell(1, column_index, str(column))
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row_index, row in enumerate(frame.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column_index, _excel_value(value))
            role = field_role(str(frame.columns[column_index - 1]), frame.iloc[:, column_index - 1])
            if role == "percentage" and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00%"
            elif role == "amount" and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
            elif role == "quantity" and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, column in enumerate(frame.columns, start=1):
        sample = [str(column), *[str(_excel_value(value) or "") for value in frame.iloc[:100, column_index - 1]]]
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(len(value) for value in sample) + 2, 36)
    return sheet


def _summary_frames(frame: pd.DataFrame, instructions: str) -> tuple[pd.DataFrame, pd.DataFrame | None, str | None]:
    overall = pd.DataFrame(summary_rows(frame))
    dimensions = []
    for column in frame.columns:
        role = field_role(str(column), frame[column])
        unique_count = frame[column].dropna().astype(str).nunique()
        if role == "category" and 1 < unique_count <= 50:
            dimensions.append(str(column))
    mentioned = [column for column in dimensions if column in instructions]
    common = [column for column in dimensions if re.search(r"渠道|地区|区域|产品|商品|客户|供应商", column)]
    dimension = (mentioned or common or [None])[0]
    if not dimension:
        return overall, None, None
    aggregations: dict[str, list[str]] = {}
    for item in summary_rows(frame)[:8]:
        aggregations[item["field"]] = ["sum", "mean"]
    if not aggregations:
        grouped = frame.groupby(dimension, dropna=False).size().reset_index(name="记录数")
    else:
        grouped = frame.groupby(dimension, dropna=False).agg(aggregations).reset_index()
        grouped.columns = [
            str(column) if not isinstance(column, tuple) else "_".join(str(part) for part in column if part)
            for column in grouped.columns
        ]
    return overall, grouped, dimension


def _write_summary_sheet(workbook: Any, result: pd.DataFrame, instructions: str) -> None:
    overall, grouped, dimension = _summary_frames(result, instructions)
    if "AI数据透视" in workbook.sheetnames:
        del workbook["AI数据透视"]
    sheet = workbook.create_sheet("AI数据透视")
    sheet["A1"] = "整体指标摘要"
    sheet["A1"].font = Font(bold=True, color="3157D5", size=14)
    if overall.empty:
        sheet["A3"] = "当前结果没有适合做合计、平均、最小或最大的数值指标。"
    else:
        labels = {"field": "字段", "role": "字段角色", "count": "有效数", "sum": "合计", "average": "平均", "min": "最小", "max": "最大"}
        for column_index, column in enumerate(overall.columns, start=1):
            cell = sheet.cell(3, column_index, labels.get(str(column), str(column)))
            cell.fill = PatternFill("solid", fgColor="3157D5")
            cell.font = Font(color="FFFFFF", bold=True)
        for row_index, row in enumerate(overall.itertuples(index=False, name=None), start=4):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row_index, column_index, _excel_value(value))
    if grouped is not None:
        start = max(7, 5 + len(overall))
        sheet.cell(start, 1, f"按{dimension}汇总")
        sheet.cell(start, 1).font = Font(bold=True, color="3157D5", size=14)
        for column_index, column in enumerate(grouped.columns, start=1):
            cell = sheet.cell(start + 2, column_index, str(column))
            cell.fill = PatternFill("solid", fgColor="3157D5")
            cell.font = Font(color="FFFFFF", bold=True)
        for row_index, row in enumerate(grouped.itertuples(index=False, name=None), start=start + 3):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row_index, column_index, _excel_value(value))
    sheet.freeze_panes = "A3"
    for column_index in range(1, max(sheet.max_column, 1) + 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = 18


def _write_notes_sheet(workbook: Any, snapshot: dict[str, Any], audit: list[dict[str, Any]], warnings: list[str]) -> None:
    if "处理说明" in workbook.sheetnames:
        del workbook["处理说明"]
    sheet = workbook.create_sheet("处理说明")
    rows = [
        ("数据任务", snapshot.get("name", "")),
        ("最终要求", snapshot.get("instructions", "")),
        ("生成时间", now_iso()),
        ("原始文件保护", "原始文件只读；本文件为派生结果，不覆盖原件。"),
        ("来源文件", "、".join(item["filename"] for item in snapshot["datasets"])),
        ("来源校验 SHA-256", "；".join(f"{item['filename']}: {item['sha256']}" for item in snapshot["datasets"])),
        ("处理步骤", json.dumps(audit, ensure_ascii=False)),
        ("提示", "；".join(warnings) or "无"),
    ]
    for row_index, (label, value) in enumerate(rows, start=1):
        sheet.cell(row_index, 1, label).font = Font(bold=True, color="3157D5")
        sheet.cell(row_index, 2, value).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 100


def _matching_columns(columns: list[str], pattern: str) -> list[str]:
    return [column for column in columns if re.search(pattern, column, re.IGNORECASE)]


def read_tables(filename: str, binary: bytes) -> dict[str, pd.DataFrame]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("仅支持 Excel（.xlsx）、CSV、TSV 和对象数组 JSON")
    try:
        if suffix == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(binary)) as workbook:
                if sum(item.file_size for item in workbook.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise ValueError("Excel 解压后数据量超过 100 MB")
            frames = pd.read_excel(io.BytesIO(binary), sheet_name=None, engine="openpyxl")
        elif suffix in {".csv", ".tsv"}:
            separator = "\t" if suffix == ".tsv" else ","
            try:
                frame = pd.read_csv(io.BytesIO(binary), sep=separator, encoding="utf-8-sig")
            except UnicodeDecodeError:
                frame = pd.read_csv(io.BytesIO(binary), sep=separator, encoding="gb18030")
            frames = {Path(filename).stem: frame}
        else:
            payload = json.loads(binary.decode("utf-8-sig"))
            if not isinstance(payload, list) or not all(isinstance(item, dict) for item in payload):
                raise ValueError("JSON 必须是对象数组")
            frames = {Path(filename).stem: pd.DataFrame(payload)}
    except ValueError:
        raise
    except Exception as error:
        raise ValueError(f"文件解析失败：{error}") from error

    normalized: dict[str, pd.DataFrame] = {}
    for sheet_name, frame in frames.items():
        if len(frame) > MAX_ROWS:
            raise ValueError(f"工作表“{sheet_name}”超过 {MAX_ROWS} 行，未进行截断")
        if len(frame.columns) > MAX_COLUMNS:
            raise ValueError(f"工作表“{sheet_name}”超过 {MAX_COLUMNS} 列，未进行截断")
        frame = frame.copy()
        frame.columns = unique_columns(list(frame.columns))
        normalized[str(sheet_name)] = frame
    if not normalized:
        raise ValueError("文件中没有可读取的数据表")
    return normalized


def _require_columns(frame: pd.DataFrame, columns: list[str]) -> None:
    missing = [column for column in columns if column not in frame.columns]
    if missing:
        raise ValueError(f"字段不存在：{'、'.join(missing)}")


def _operand(frame: pd.DataFrame, value: Any) -> Any:
    if isinstance(value, dict) and value.get("column"):
        column = str(value["column"])
        _require_columns(frame, [column])
        return pd.to_numeric(frame[column], errors="coerce")
    if isinstance(value, dict) and "value" in value:
        return value["value"]
    raise ValueError("计算列操作数必须指定 column 或 value")


def apply_plan(
    frames: dict[str, pd.DataFrame],
    plan: dict[str, Any],
) -> tuple[pd.DataFrame, list[dict[str, Any]], list[str]]:
    source_ids = plan.get("sourceTableIds")
    if not isinstance(source_ids, list) or not source_ids:
        raise ValueError("处理计划必须选择至少一个数据表")
    missing_sources = [table_id for table_id in source_ids if table_id not in frames]
    if missing_sources:
        raise ValueError(f"数据表不存在：{'、'.join(missing_sources)}")

    combine_mode = plan.get("combineMode", "first")
    if combine_mode == "concat":
        frame = pd.concat([frames[table_id] for table_id in source_ids], ignore_index=True, sort=False)
    elif combine_mode == "first":
        frame = frames[source_ids[0]].copy()
    else:
        raise ValueError("combineMode 仅支持 first 或 concat")

    audit: list[dict[str, Any]] = []
    warnings: list[str] = []
    operations = plan.get("operations", [])
    if not isinstance(operations, list):
        raise ValueError("operations 必须是数组")

    for index, operation in enumerate(operations, start=1):
        if not isinstance(operation, dict):
            raise ValueError(f"第 {index} 个操作格式无效")
        kind = str(operation.get("type", ""))
        before_rows, before_columns = len(frame), len(frame.columns)

        if kind == "select_columns":
            columns = [str(item) for item in operation.get("columns", [])]
            _require_columns(frame, columns)
            frame = frame.loc[:, columns].copy()
        elif kind == "rename_columns":
            mapping = operation.get("mapping", {})
            if not isinstance(mapping, dict) or not mapping:
                raise ValueError("重命名必须提供 mapping")
            _require_columns(frame, [str(column) for column in mapping])
            target_names = [str(value).strip() for value in mapping.values()]
            if any(not value for value in target_names) or len(target_names) != len(set(target_names)):
                raise ValueError("重命名后的字段不能为空或重复")
            frame = frame.rename(columns={str(key): str(value).strip() for key, value in mapping.items()})
            if len(set(frame.columns)) != len(frame.columns):
                raise ValueError("重命名后产生了重复字段")
        elif kind == "reorder_columns":
            columns = [str(item) for item in operation.get("columns", [])]
            _require_columns(frame, columns)
            remaining = [column for column in frame.columns if column not in columns]
            frame = frame.loc[:, columns + remaining].copy()
        elif kind == "trim":
            columns = [str(item) for item in operation.get("columns", [])] or [
                str(column) for column in frame.select_dtypes(include=["object", "string"]).columns
            ]
            _require_columns(frame, columns)
            for column in columns:
                frame[column] = frame[column].map(lambda value: value.strip() if isinstance(value, str) else value)
        elif kind == "normalize_empty":
            columns = [str(item) for item in operation.get("columns", [])] or [str(column) for column in frame.columns]
            markers = operation.get("markers", ["", "-", "--", "null", "none", "n/a", "na"])
            normalized_markers = {str(item).strip().lower() for item in markers}
            _require_columns(frame, columns)
            for column in columns:
                frame[column] = frame[column].map(
                    lambda value: pd.NA
                    if value is None or (isinstance(value, str) and value.strip().lower() in normalized_markers)
                    else value
                )
        elif kind == "convert_type":
            column = str(operation.get("column", ""))
            target = str(operation.get("targetType", ""))
            _require_columns(frame, [column])
            if target == "number":
                frame[column] = pd.to_numeric(frame[column], errors="coerce")
            elif target == "date":
                frame[column] = pd.to_datetime(frame[column], errors="coerce")
            elif target == "text":
                frame[column] = frame[column].astype("string")
            elif target == "boolean":
                truthy = {"1", "true", "yes", "是", "对"}
                falsy = {"0", "false", "no", "否", "错"}
                def boolean_value(value: Any) -> Any:
                    normalized = str(value).strip().lower()
                    if normalized in truthy:
                        return True
                    if normalized in falsy:
                        return False
                    return pd.NA
                frame[column] = frame[column].map(boolean_value).astype("boolean")
            else:
                raise ValueError("目标类型仅支持 number、date、text、boolean")
        elif kind == "drop_duplicates":
            columns = [str(item) for item in operation.get("columns", [])]
            if columns:
                _require_columns(frame, columns)
            frame = frame.drop_duplicates(subset=columns or None).reset_index(drop=True)
        elif kind == "filter":
            column = str(operation.get("column", ""))
            operator = str(operation.get("operator", "eq"))
            value = operation.get("value")
            _require_columns(frame, [column])
            series = frame[column]
            if operator in {"gt", "gte", "lt", "lte"}:
                left = pd.to_numeric(series, errors="coerce")
                try:
                    right = float(value)
                except (TypeError, ValueError) as error:
                    raise ValueError("数值筛选条件无效") from error
                masks = {"gt": left > right, "gte": left >= right, "lt": left < right, "lte": left <= right}
                mask = masks[operator]
            elif operator == "contains":
                mask = series.astype("string").str.contains(str(value), case=False, na=False, regex=False)
            elif operator == "starts_with":
                mask = series.astype("string").str.startswith(str(value), na=False)
            elif operator == "is_empty":
                mask = series.isna() | (series.astype("string").str.strip() == "")
            elif operator == "not_empty":
                mask = series.notna() & (series.astype("string").str.strip() != "")
            elif operator == "ne":
                mask = series.astype("string") != str(value)
            elif operator == "eq":
                mask = series.astype("string") == str(value)
            else:
                raise ValueError("不支持的筛选条件")
            frame = frame.loc[mask.fillna(False)].reset_index(drop=True)
        elif kind == "sort":
            columns_config = operation.get("columns", [])
            columns = [str(item.get("column", "")) for item in columns_config]
            _require_columns(frame, columns)
            frame = frame.sort_values(
                by=columns,
                ascending=[bool(item.get("ascending", True)) for item in columns_config],
                na_position="last",
            ).reset_index(drop=True)
        elif kind == "split_column":
            column = str(operation.get("column", ""))
            separator = str(operation.get("separator", ""))
            new_columns = [str(item) for item in operation.get("newColumns", [])]
            _require_columns(frame, [column])
            if not separator or len(new_columns) < 2:
                raise ValueError("拆分字段需要分隔符和至少两个新字段")
            split = frame[column].astype("string").str.split(separator, n=len(new_columns) - 1, expand=True)
            for position, new_column in enumerate(new_columns):
                frame[new_column] = split[position] if position in split.columns else pd.NA
        elif kind == "merge_columns":
            columns = [str(item) for item in operation.get("columns", [])]
            new_column = str(operation.get("newColumn", "")).strip()
            separator = str(operation.get("separator", ""))
            _require_columns(frame, columns)
            if len(columns) < 2 or not new_column:
                raise ValueError("合并字段需要至少两个来源字段和新字段名")
            frame[new_column] = frame[columns].fillna("").astype(str).agg(separator.join, axis=1)
        elif kind == "calculate":
            new_column = str(operation.get("newColumn", "")).strip()
            operator = str(operation.get("operator", ""))
            if not new_column:
                raise ValueError("计算列必须提供新字段名")
            left = _operand(frame, operation.get("left"))
            right = _operand(frame, operation.get("right"))
            if operator == "add":
                frame[new_column] = left + right
            elif operator == "subtract":
                frame[new_column] = left - right
            elif operator == "multiply":
                frame[new_column] = left * right
            elif operator == "divide":
                frame[new_column] = left / right
                frame[new_column] = frame[new_column].replace([math.inf, -math.inf], pd.NA)
            else:
                raise ValueError("计算列仅支持 add、subtract、multiply、divide")
        elif kind == "group":
            by = [str(item) for item in operation.get("by", [])]
            aggregations = operation.get("aggregations", [])
            _require_columns(frame, by + [str(item.get("column", "")) for item in aggregations])
            if not by or not aggregations:
                raise ValueError("分组汇总必须提供分组字段和汇总项")
            named: dict[str, pd.NamedAgg] = {}
            allowed = {"sum", "mean", "min", "max", "count", "nunique"}
            for item in aggregations:
                operation_name = str(item.get("operation", ""))
                if operation_name not in allowed:
                    raise ValueError("汇总方式仅支持 sum、mean、min、max、count、nunique")
                source = str(item["column"])
                output = str(item.get("outputColumn") or f"{source}_{operation_name}")
                named[output] = pd.NamedAgg(column=source, aggfunc=operation_name)
            frame = frame.groupby(by, dropna=False).agg(**named).reset_index()
        elif kind == "join":
            right_id = str(operation.get("rightTableId", ""))
            if right_id not in frames:
                raise ValueError("关联的右侧数据表不存在")
            left_on = [str(item) for item in operation.get("leftOn", [])]
            right_on = [str(item) for item in operation.get("rightOn", [])]
            how = str(operation.get("how", "left"))
            if how not in {"left", "inner", "outer", "right"} or not left_on or len(left_on) != len(right_on):
                raise ValueError("关联方式或关联字段无效")
            right_frame = frames[right_id]
            _require_columns(frame, left_on)
            _require_columns(right_frame, right_on)
            if right_frame.duplicated(subset=right_on).any():
                warnings.append(f"右表 {right_id} 的关联键存在重复，结果可能产生一对多行")
            frame = frame.merge(right_frame, how=how, left_on=left_on, right_on=right_on, suffixes=("", "_右表"))
        else:
            raise ValueError(f"不支持的操作类型：{kind or '空'}")

        if len(frame) > MAX_ROWS:
            raise ValueError(f"第 {index} 步执行后超过 {MAX_ROWS} 行，已阻断且未截断")
        audit.append({
            "step": index,
            "type": kind,
            "beforeRows": int(before_rows),
            "afterRows": int(len(frame)),
            "beforeColumns": int(before_columns),
            "afterColumns": int(len(frame.columns)),
        })

    return frame, audit, warnings


class DataWorkspaceStore:
    def __init__(self, root: Path) -> None:
        self.root = root
        self.root.mkdir(parents=True, exist_ok=True)
        self.metadata_file = root / "workspaces.json"
        self._lock = Lock()
        self._records: dict[str, dict[str, Any]] = {}
        if self.metadata_file.exists():
            payload = json.loads(self.metadata_file.read_text(encoding="utf-8"))
            self._records = payload.get("workspaces", {})

    def _save_locked(self) -> None:
        temporary = self.metadata_file.with_suffix(".tmp")
        temporary.write_text(json.dumps({"workspaces": self._records}, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(self.metadata_file)

    def create(self, name: str, instructions: str = "") -> dict[str, Any]:
        workspace_id = uuid.uuid4().hex
        record = {
            "id": workspace_id,
            "name": name.strip() or "未命名数据任务",
            "instructions": instructions.strip(),
            "status": "waiting_for_data",
            "datasets": [],
            "tables": [],
            "plan": None,
            "planVersion": 0,
            "preview": None,
            "execution": None,
            "timeline": [{"stage": "workspace_created", "at": now_iso()}],
            "createdAt": now_iso(),
            "updatedAt": now_iso(),
        }
        with self._lock:
            self._records[workspace_id] = record
            self._save_locked()
        (self.root / workspace_id / "sources").mkdir(parents=True, exist_ok=True)
        (self.root / workspace_id / "outputs").mkdir(parents=True, exist_ok=True)
        return dict(record)

    def list(self) -> list[dict[str, Any]]:
        with self._lock:
            return [dict(item) for item in sorted(self._records.values(), key=lambda item: item["createdAt"], reverse=True)]

    def get(self, workspace_id: str) -> dict[str, Any] | None:
        with self._lock:
            record = self._records.get(workspace_id)
            return dict(record) if record else None

    def delete(self, workspace_id: str) -> dict[str, Any]:
        with self._lock:
            if workspace_id not in self._records:
                raise KeyError("数据任务不存在")
        workspace_dir = self.root / workspace_id
        if workspace_dir.is_dir():
            shutil.rmtree(workspace_dir)
        with self._lock:
            if workspace_id not in self._records:
                raise KeyError("数据任务不存在")
            del self._records[workspace_id]
            self._save_locked()
        return {"deleted": workspace_id}

    def upload(self, workspace_id: str, filename: str, content_base64: str) -> dict[str, Any]:
        safe_name = Path(filename).name
        if Path(safe_name).suffix.lower() not in SUPPORTED_SUFFIXES:
            raise ValueError("不支持的文件类型")
        try:
            binary = base64.b64decode(content_base64, validate=True)
        except Exception as error:
            raise ValueError("文件编码无效") from error
        if not binary:
            raise ValueError("文件不能为空")
        if len(binary) > MAX_FILE_BYTES:
            raise ValueError("单文件不能超过 20 MB")

        with self._lock:
            record = self._records.get(workspace_id)
            if not record:
                raise KeyError("数据任务不存在")
            if len(record["datasets"]) >= MAX_FILES:
                raise ValueError("单个任务最多上传 5 个文件")

        tables = read_tables(safe_name, binary)
        dataset_id = uuid.uuid4().hex[:12]
        source_name = f"{dataset_id}-{safe_name}"
        source_path = self.root / workspace_id / "sources" / source_name
        source_path.write_bytes(binary)
        os.chmod(source_path, 0o444)
        digest = hashlib.sha256(binary).hexdigest()
        table_records = []
        for sheet_name, frame in tables.items():
            table_id = f"{dataset_id}:{sheet_name}"
            table_records.append({
                "id": table_id,
                "datasetId": dataset_id,
                "name": sheet_name,
                "profile": profile_frame(frame),
                "sample": frame_records(frame),
            })
        dataset = {
            "id": dataset_id,
            "filename": safe_name,
            "storedName": source_name,
            "sha256": digest,
            "size": len(binary),
            "uploadedAt": now_iso(),
        }
        with self._lock:
            record = self._records.get(workspace_id)
            if not record:
                raise KeyError("数据任务不存在")
            record["datasets"].append(dataset)
            record["tables"].extend(table_records)
            record["status"] = "plan_draft"
            record["preview"] = None
            record["execution"] = None
            record["updatedAt"] = now_iso()
            record["timeline"].append({"stage": "dataset_uploaded", "at": now_iso(), "datasetId": dataset_id})
            self._save_locked()
            return dict(record)

    def _frames(self, record: dict[str, Any]) -> dict[str, pd.DataFrame]:
        frames: dict[str, pd.DataFrame] = {}
        for dataset in record["datasets"]:
            source_path = self.root / record["id"] / "sources" / dataset["storedName"]
            binary = source_path.read_bytes()
            if hashlib.sha256(binary).hexdigest() != dataset["sha256"]:
                raise ValueError(f"原始文件校验失败：{dataset['filename']}")
            for sheet_name, frame in read_tables(dataset["filename"], binary).items():
                frames[f"{dataset['id']}:{sheet_name}"] = frame
        return frames

    def suggest_plan(self, workspace_id: str, instructions: str) -> dict[str, Any]:
        with self._lock:
            record = self._records.get(workspace_id)
            if not record:
                raise KeyError("数据任务不存在")
            if not record["tables"]:
                raise ValueError("请先上传数据")
            table_id = record["tables"][0]["id"]
        text = instructions.strip()
        operations: list[dict[str, Any]] = []
        if re.search(r"空格|整理|清理|清洗", text):
            operations.append({"type": "trim", "columns": []})
        if re.search(r"空值|空白|null|缺失", text):
            operations.append({"type": "normalize_empty", "columns": []})
        if re.search(r"去重|重复", text):
            operations.append({"type": "drop_duplicates", "columns": []})
        if not operations:
            operations = [{"type": "trim", "columns": []}, {"type": "normalize_empty", "columns": []}]
        return {
            "instructions": text,
            "sourceTableIds": [table_id],
            "combineMode": "first",
            "operations": operations,
            "warnings": ["系统只生成保守建议；筛选、匹配、删除和计算口径需在操作计划中明确确认。"],
        }

    def infer_plan(
        self,
        workspace_id: str,
        instructions: str,
        decisions: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        with self._lock:
            record = self._records.get(workspace_id)
            if not record:
                raise KeyError("数据任务不存在")
            if not record["tables"]:
                raise ValueError("请先上传数据")
            tables = json.loads(json.dumps(record["tables"]))
        text = instructions.strip()
        if not text:
            raise ValueError("请用一句话说明最终想要的结果")
        decisions = decisions or {}
        source_ids = [item["id"] for item in tables]
        combine_mode = "first"
        operations: list[dict[str, Any]] = [
            {"type": "trim", "columns": []},
            {"type": "normalize_empty", "columns": []},
        ]
        if re.search(r"去重|删除重复", text):
            operations.append({"type": "drop_duplicates", "columns": []})
        if len(source_ids) > 1 and re.search(r"纵向合并|追加|合并同类", text):
            first_columns = tables[0]["profile"]["columns"]
            if all(item["profile"]["columns"] == first_columns for item in tables[1:]):
                combine_mode = "concat"
            else:
                return {
                    "status": "needs_input",
                    "resolution": {
                        "type": "information",
                        "field": "combineMode",
                        "title": "多张表字段不一致",
                        "reason": "直接纵向合并会产生错列或大量空值，需要先在高级设置中明确匹配方式。",
                        "choices": [{"value": "advanced", "label": "打开高级设置"}],
                    },
                }
        columns = tables[0]["profile"]["columns"]
        if "利润" in text:
            sales = _matching_columns(columns, r"销售额|销售收入|航司售价|售价|收入|金额")
            costs = _matching_columns(columns, r"结算底价|修改底价|底价|采购成本|成本|进价")
            mentioned_sales = [column for column in sales if column in text]
            mentioned_costs = [column for column in costs if column in text]
            sales_choice = str(decisions.get("salesColumn", "")) or (mentioned_sales[0] if len(mentioned_sales) == 1 else "")
            cost_choice = str(decisions.get("costColumn", "")) or (mentioned_costs[0] if len(mentioned_costs) == 1 else "")
            if not sales_choice and len(sales) == 1:
                sales_choice = sales[0]
            if not cost_choice and len(costs) == 1:
                cost_choice = costs[0]
            for field, title, reason, choices in (
                ("salesColumn", "请选择收入字段", "系统找到多个可能的收入字段，直接选择会改变利润结果。", sales),
                ("costColumn", "请选择成本字段", "系统找到多个可能的成本字段，直接选择会改变利润结果。", costs),
            ):
                current = sales_choice if field == "salesColumn" else cost_choice
                if not current:
                    if not choices:
                        raise ValueError(f"无法计算利润：没有找到{title.replace('请选择', '')}")
                    return {
                        "status": "needs_input",
                        "resolution": {
                            "type": "column_choice",
                            "field": field,
                            "title": title,
                            "reason": reason,
                            "choices": [{"value": column, "label": column} for column in choices],
                            "decisions": decisions,
                        },
                    }
            if sales_choice not in columns or cost_choice not in columns:
                raise ValueError("选择的利润计算字段不存在")
            operations.append({
                "type": "calculate",
                "newColumn": "利润",
                "left": {"column": sales_choice},
                "operator": "subtract",
                "right": {"column": cost_choice},
            })
        return {
            "status": "ready",
            "plan": {
                "instructions": text,
                "sourceTableIds": source_ids if combine_mode == "concat" else [source_ids[0]],
                "combineMode": combine_mode,
                "operations": operations,
            },
        }

    def save_plan(self, workspace_id: str, plan: dict[str, Any]) -> dict[str, Any]:
        with self._lock:
            record = self._records.get(workspace_id)
            if not record:
                raise KeyError("数据任务不存在")
        frames = self._frames(record)
        apply_plan(frames, plan)
        with self._lock:
            record = self._records[workspace_id]
            record["planVersion"] += 1
            record["plan"] = plan
            record["instructions"] = str(plan.get("instructions", record.get("instructions", ""))).strip()
            record["preview"] = None
            record["execution"] = None
            record["status"] = "plan_saved"
            record["updatedAt"] = now_iso()
            record["timeline"].append({"stage": "plan_saved", "at": now_iso(), "version": record["planVersion"]})
            self._save_locked()
            return dict(record)

    def preview(self, workspace_id: str) -> dict[str, Any]:
        with self._lock:
            record = self._records.get(workspace_id)
            if not record:
                raise KeyError("数据任务不存在")
            if not record.get("plan"):
                raise ValueError("请先保存处理计划")
            snapshot = json.loads(json.dumps(record))
        frames = self._frames(snapshot)
        source_id = snapshot["plan"]["sourceTableIds"][0]
        result, audit, warnings = apply_plan(frames, snapshot["plan"])
        token_payload = json.dumps({
            "workspace": workspace_id,
            "version": snapshot["planVersion"],
            "plan": snapshot["plan"],
            "sources": [item["sha256"] for item in snapshot["datasets"]],
        }, ensure_ascii=False, sort_keys=True).encode("utf-8")
        preview_id = hashlib.sha256(token_payload).hexdigest()
        preview = {
            "id": preview_id,
            "planVersion": snapshot["planVersion"],
            "source": {"profile": profile_frame(frames[source_id]), "rows": frame_records(frames[source_id])},
            "result": {"profile": profile_frame(result), "rows": frame_records(result)},
            "audit": audit,
            "warnings": warnings,
            "generatedAt": now_iso(),
        }
        with self._lock:
            record = self._records[workspace_id]
            record["preview"] = preview
            record["status"] = "preview_ready"
            record["updatedAt"] = now_iso()
            record["timeline"].append({"stage": "preview_generated", "at": now_iso(), "previewId": preview_id})
            self._save_locked()
        return preview

    def execute(self, workspace_id: str, preview_id: str, output_modes: list[str] | None = None) -> dict[str, Any]:
        with self._lock:
            record = self._records.get(workspace_id)
            if not record:
                raise KeyError("数据任务不存在")
            if not record.get("preview") or record["preview"].get("id") != preview_id:
                raise ValueError("预览已失效，请重新预览并确认")
            snapshot = json.loads(json.dumps(record))
        frames = self._frames(snapshot)
        result, audit, warnings = apply_plan(frames, snapshot["plan"])
        output_dir = self.root / workspace_id / "outputs"
        selected_modes = list(dict.fromkeys(output_modes or ["standalone"]))
        if not selected_modes or any(mode not in OUTPUT_MODES for mode in selected_modes):
            raise ValueError("导出方式仅支持关联原表或独立结果")
        exports: dict[str, dict[str, str]] = {}
        if "standalone" in selected_modes:
            csv_path = output_dir / "result.csv"
            xlsx_path = output_dir / "result.xlsx"
            result.to_csv(csv_path, index=False, encoding="utf-8-sig")
            workbook = Workbook()
            del workbook[workbook.sheetnames[0]]
            _write_frame_sheet(workbook, "AI处理结果", result)
            _write_summary_sheet(workbook, result, snapshot.get("instructions", ""))
            _write_notes_sheet(workbook, snapshot, audit, warnings)
            workbook.save(xlsx_path)
            exports.update({
                "csv": {"filename": "result.csv", "sha256": hashlib.sha256(csv_path.read_bytes()).hexdigest()},
                "xlsx": {"filename": "result.xlsx", "sha256": hashlib.sha256(xlsx_path.read_bytes()).hexdigest()},
            })
        if "linked" in selected_modes:
            linked_path = output_dir / "linked-result.xlsx"
            first_dataset = snapshot["datasets"][0]
            source_path = self.root / workspace_id / "sources" / first_dataset["storedName"]
            if Path(first_dataset["filename"]).suffix.lower() == ".xlsx":
                shutil.copy2(source_path, linked_path)
                os.chmod(linked_path, 0o644)
                workbook = load_workbook(linked_path)
            else:
                workbook = Workbook()
                del workbook[workbook.sheetnames[0]]
                for table_id, frame in frames.items():
                    title = table_id.split(":", 1)[-1][:31]
                    _write_frame_sheet(workbook, title, frame)
            _write_frame_sheet(workbook, "AI处理结果", result)
            _write_summary_sheet(workbook, result, snapshot.get("instructions", ""))
            _write_notes_sheet(workbook, snapshot, audit, warnings)
            workbook.save(linked_path)
            exports["linked"] = {"filename": "linked-result.xlsx", "sha256": hashlib.sha256(linked_path.read_bytes()).hexdigest()}
        execution = {
            "previewId": preview_id,
            "planVersion": snapshot["planVersion"],
            "profile": profile_frame(result),
            "audit": audit,
            "warnings": warnings,
            "summary": summary_rows(result),
            "outputModes": selected_modes,
            "exports": exports,
            "completedAt": now_iso(),
        }
        with self._lock:
            record = self._records[workspace_id]
            record["execution"] = execution
            record["status"] = "completed"
            record["updatedAt"] = now_iso()
            record["timeline"].append({"stage": "plan_executed", "at": now_iso(), "previewId": preview_id})
            self._save_locked()
        return execution

    def auto_run(
        self,
        workspace_id: str,
        instructions: str,
        output_modes: list[str] | None = None,
        decisions: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        selected_modes = output_modes or ["linked"]
        inferred = self.infer_plan(workspace_id, instructions, decisions)
        if inferred["status"] == "needs_input":
            with self._lock:
                record = self._records[workspace_id]
                record["instructions"] = instructions.strip()
                record["status"] = "needs_input"
                record["pendingResolution"] = inferred["resolution"]
                record["outputModes"] = selected_modes
                record["updatedAt"] = now_iso()
                record["timeline"].append({"stage": "input_needed", "at": now_iso(), "field": inferred["resolution"]["field"]})
                self._save_locked()
                return {"status": "needs_input", "workspace": dict(record), "resolution": inferred["resolution"]}
        self.save_plan(workspace_id, inferred["plan"])
        preview = self.preview(workspace_id)
        execution = self.execute(workspace_id, preview["id"], selected_modes)
        with self._lock:
            record = self._records[workspace_id]
            record.pop("pendingResolution", None)
            record["outputModes"] = selected_modes
            self._save_locked()
            workspace = dict(record)
        return {"status": "completed", "workspace": workspace, "execution": execution}

    def export_path(self, workspace_id: str, file_format: str) -> Path:
        filenames = {"csv": "result.csv", "xlsx": "result.xlsx", "linked": "linked-result.xlsx"}
        if file_format not in filenames:
            raise ValueError("导出格式仅支持 csv、xlsx 或 linked")
        with self._lock:
            record = self._records.get(workspace_id)
            if not record or not record.get("execution"):
                raise ValueError("任务尚未执行，不能导出")
        path = self.root / workspace_id / "outputs" / filenames[file_format]
        if not path.is_file():
            raise ValueError("导出文件不存在")
        return path
