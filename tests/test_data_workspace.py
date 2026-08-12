import base64
import hashlib
import io
import json
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from data_engine import DataWorkspaceStore, apply_plan, field_role, read_tables, summary_rows


def encoded_csv(text: str) -> str:
    return base64.b64encode(text.encode("utf-8-sig")).decode("ascii")


class DataWorkspaceTest(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.store = DataWorkspaceStore(Path(self.temporary_directory.name) / "workspaces")

    def tearDown(self):
        self.temporary_directory.cleanup()

    def create_with_sales(self):
        workspace = self.store.create("销售数据整理", "清理空格、去重并筛选")
        workspace = self.store.upload(
            workspace["id"],
            "sales.csv",
            encoded_csv("渠道,销售额,成本\n 官网 ,100,60\n官网,100,60\n门店,200,150\n批发,,30\n"),
        )
        return workspace

    def test_upload_profiles_data_and_keeps_source_read_only(self):
        workspace = self.create_with_sales()
        dataset = workspace["datasets"][0]
        source = Path(self.temporary_directory.name) / "workspaces" / workspace["id"] / "sources" / dataset["storedName"]

        self.assertEqual(workspace["tables"][0]["profile"]["rowCount"], 4)
        self.assertEqual(workspace["tables"][0]["profile"]["columns"], ["渠道", "销售额", "成本"])
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), dataset["sha256"])
        self.assertEqual(source.stat().st_mode & 0o222, 0)

    def test_editable_plan_previews_then_executes_without_touching_source(self):
        workspace = self.create_with_sales()
        dataset = workspace["datasets"][0]
        source = Path(self.temporary_directory.name) / "workspaces" / workspace["id"] / "sources" / dataset["storedName"]
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        table_id = workspace["tables"][0]["id"]
        plan = {
            "sourceTableIds": [table_id],
            "combineMode": "first",
            "operations": [
                {"type": "trim", "columns": ["渠道"]},
                {"type": "drop_duplicates", "columns": ["渠道", "销售额", "成本"]},
                {"type": "filter", "column": "销售额", "operator": "gte", "value": 100},
                {
                    "type": "calculate",
                    "newColumn": "利润",
                    "left": {"column": "销售额"},
                    "operator": "subtract",
                    "right": {"column": "成本"},
                },
                {"type": "sort", "columns": [{"column": "利润", "ascending": False}]},
            ],
        }
        self.store.save_plan(workspace["id"], plan)
        preview = self.store.preview(workspace["id"])

        self.assertEqual(preview["source"]["profile"]["rowCount"], 4)
        self.assertEqual(preview["result"]["profile"]["rowCount"], 2)
        self.assertEqual(preview["result"]["rows"][0]["利润"], 50.0)
        self.assertEqual(self.store.get(workspace["id"])["status"], "preview_ready")
        with self.assertRaisesRegex(ValueError, "预览已失效"):
            self.store.execute(workspace["id"], "wrong-preview")

        execution = self.store.execute(workspace["id"], preview["id"])
        self.assertEqual(execution["profile"]["rowCount"], 2)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)
        self.assertTrue(self.store.export_path(workspace["id"], "csv").is_file())
        self.assertTrue(self.store.export_path(workspace["id"], "xlsx").is_file())

    def test_operations_cover_rename_reorder_split_merge_group_and_type_conversion(self):
        frame = pd.DataFrame({
            "地区渠道": ["华东-线上", "华东-线下", "华北-线上"],
            "金额": ["10", "20", "30"],
        })
        result, audit, _ = apply_plan(
            {"source": frame},
            {
                "sourceTableIds": ["source"],
                "combineMode": "first",
                "operations": [
                    {"type": "rename_columns", "mapping": {"金额": "销售额"}},
                    {"type": "convert_type", "column": "销售额", "targetType": "number"},
                    {"type": "split_column", "column": "地区渠道", "separator": "-", "newColumns": ["地区", "渠道"]},
                    {"type": "merge_columns", "columns": ["地区", "渠道"], "newColumn": "业务范围", "separator": "/"},
                    {"type": "group", "by": ["地区"], "aggregations": [{"column": "销售额", "operation": "sum", "outputColumn": "销售合计"}]},
                    {"type": "reorder_columns", "columns": ["销售合计", "地区"]},
                ],
            },
        )

        self.assertEqual(list(result.columns), ["销售合计", "地区"])
        self.assertEqual(result.loc[result["地区"] == "华东", "销售合计"].iloc[0], 30)
        self.assertEqual(len(audit), 6)

    def test_multiple_files_can_concat_and_join_with_duplicate_key_warning(self):
        frames = {
            "january": pd.DataFrame({"订单号": [1, 2], "金额": [10, 20]}),
            "february": pd.DataFrame({"订单号": [3], "金额": [30]}),
            "customers": pd.DataFrame({"订单号": [1, 1, 3], "客户": ["甲", "甲重复", "乙"]}),
        }
        result, _, warnings = apply_plan(
            frames,
            {
                "sourceTableIds": ["january", "february"],
                "combineMode": "concat",
                "operations": [{
                    "type": "join",
                    "rightTableId": "customers",
                    "leftOn": ["订单号"],
                    "rightOn": ["订单号"],
                    "how": "left",
                }],
            },
        )

        self.assertEqual(len(result), 4)
        self.assertTrue(any("一对多" in warning for warning in warnings))

    def test_xlsx_reads_all_sheets_and_exports_valid_excel(self):
        binary = io.BytesIO()
        with pd.ExcelWriter(binary, engine="openpyxl") as writer:
            pd.DataFrame({"订单号": [1], "金额": [10]}).to_excel(writer, sheet_name="订单", index=False)
            pd.DataFrame({"订单号": [1], "状态": ["已付"]}).to_excel(writer, sheet_name="支付", index=False)

        tables = read_tables("business.xlsx", binary.getvalue())

        self.assertEqual(set(tables), {"订单", "支付"})
        self.assertEqual(tables["订单"].iloc[0]["金额"], 10)

    def test_wide_data_is_rejected_instead_of_silently_truncated(self):
        header = ",".join(f"字段{index}" for index in range(1001))
        values = ",".join("1" for _ in range(1001))

        with self.assertRaisesRegex(ValueError, "超过 1000 列"):
            read_tables("wide.csv", f"{header}\n{values}\n".encode("utf-8"))

    def test_workspace_persists_after_store_restart(self):
        workspace = self.create_with_sales()
        reloaded = DataWorkspaceStore(Path(self.temporary_directory.name) / "workspaces")

        self.assertEqual(reloaded.get(workspace["id"])["datasets"][0]["filename"], "sales.csv")

    def test_workspace_delete_removes_record_and_local_files(self):
        workspace = self.create_with_sales()
        workspace_dir = Path(self.temporary_directory.name) / "workspaces" / workspace["id"]

        result = self.store.delete(workspace["id"])

        self.assertEqual(result["deleted"], workspace["id"])
        self.assertIsNone(self.store.get(workspace["id"]))
        self.assertFalse(workspace_dir.exists())

    def test_suggested_plan_is_conservative_and_editable(self):
        workspace = self.create_with_sales()
        suggestion = self.store.suggest_plan(workspace["id"], "清理空格、统一空值并去重")

        self.assertEqual(
            [operation["type"] for operation in suggestion["operations"]],
            ["trim", "normalize_empty", "drop_duplicates"],
        )
        self.assertTrue(suggestion["warnings"])

    def test_semantic_roles_do_not_sum_identifiers_or_dates(self):
        frame = pd.DataFrame({"出票单号": [68837839861985], "出票日期": [45868], "扣率": [0.95], "航司售价": [600]})

        self.assertEqual(field_role("出票单号", frame["出票单号"]), "identifier")
        self.assertEqual(field_role("出票日期", frame["出票日期"]), "date")
        self.assertEqual(field_role("扣率", frame["扣率"]), "percentage")
        self.assertEqual([item["field"] for item in summary_rows(frame)], ["扣率", "航司售价"])

    def test_one_click_linked_export_preserves_original_and_adds_ai_sheets(self):
        source = io.BytesIO()
        with pd.ExcelWriter(source, engine="openpyxl") as writer:
            pd.DataFrame({"渠道": ["官网", "门店"], "航司售价": [100, 200], "结算底价": [60, 150]}).to_excel(writer, sheet_name="原始明细", index=False)
        workspace = self.store.create("利润", "计算利润")
        workspace = self.store.upload(workspace["id"], "业务.xlsx", base64.b64encode(source.getvalue()).decode("ascii"))
        source_file = Path(self.temporary_directory.name) / "workspaces" / workspace["id"] / "sources" / workspace["datasets"][0]["storedName"]
        original_hash = hashlib.sha256(source_file.read_bytes()).hexdigest()

        result = self.store.auto_run(workspace["id"], "计算航司售价减结算底价后的利润，按渠道汇总", ["linked", "standalone"])

        self.assertEqual(result["status"], "completed")
        linked = self.store.export_path(workspace["id"], "linked")
        workbook = load_workbook(linked, read_only=True)
        self.assertTrue({"原始明细", "AI处理结果", "AI数据透视", "处理说明"}.issubset(workbook.sheetnames))
        self.assertEqual(hashlib.sha256(source_file.read_bytes()).hexdigest(), original_hash)
        self.assertEqual(result["execution"]["profile"]["rowCount"], 2)

    def test_one_click_explains_ambiguous_cost_and_resumes_after_choice(self):
        workspace = self.store.create("利润歧义", "计算利润")
        workspace = self.store.upload(
            workspace["id"],
            "sales.csv",
            encoded_csv("航司售价,结算底价,修改底价\n100,60,55\n"),
        )

        blocked = self.store.auto_run(workspace["id"], "计算利润", ["standalone"])
        self.assertEqual(blocked["status"], "needs_input")
        self.assertEqual(blocked["resolution"]["field"], "costColumn")
        completed = self.store.auto_run(
            workspace["id"], "计算利润", ["standalone"], {"costColumn": "结算底价"}
        )
        self.assertEqual(completed["status"], "completed")
        self.assertEqual(completed["workspace"]["preview"]["result"]["rows"][0]["利润"], 40)


if __name__ == "__main__":
    unittest.main()
